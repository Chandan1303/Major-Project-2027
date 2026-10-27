"""Authenticated API for variety-specific crop growth monitoring.

This blueprint is independent of the yield prediction service and reuses the
existing authenticated user, farm, and field records.
"""

import csv
import hashlib
import io
import json
import math
import re
import uuid
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import jwt
from flask import Blueprint, current_app, jsonify, request, send_file, send_from_directory
from openpyxl import load_workbook
from PIL import Image
from werkzeug.utils import secure_filename

from backend.crop_intelligence_models import (
    CropActivity, CropAlert, CropAlertEvent, CropComparison, CropCycle, CropImport,
    CropMitigationStrategy, CropObservation, CropObservationRevision,
    CropProgressReport, CropRecommendation, CropStageReference, CropTask, CropTaskEvent,
    CropVarietyReference,
)
from backend.models import Farm, Field, User, db
from backend.crop_intelligence_service import (
    CHECKPOINTS, DATA_DIR, MITIGATION_SEED, NUMERIC_METRICS, OBSERVATION_FIELDS, STAGE_SEED,
    VARIETY_SEED, alert_dict, assess_observation, dap_for, expected_stage_guidance,
    json_text, next_checkpoint, observation_dict, parse_date, parse_json,
    report_payload, stage_for_calendar_day, stage_for_dap, stage_progress,
    summarize_stage, task_dict,
)


crop_intelligence_bp = Blueprint("crop_intelligence", __name__, url_prefix="/api/crop-intelligence")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
OBSERVATION_KEYS = set(OBSERVATION_FIELDS)
NUMERIC_LIMITS = {
    "plant_height_cm": (0, 2000), "tiller_count": (0, 1_000_000),
    "measurement_area_m2": (0, 10_000_000), "establishment_pct": (0, 100),
    "soil_moisture_pct": (0, 100), "rainfall_mm": (0, 100_000),
}
LEAF_CONDITIONS = {"normal", "yellowing", "drying", "curling", "discoloration", "wilting", "other"}
TASK_STATUSES = {"pending", "in_progress", "completed", "cancelled"}
ALERT_STATUSES = {"new", "acknowledged", "monitoring", "resolved", "unconfirmed"}


@crop_intelligence_bp.before_request
def authenticate_crop_intelligence():
    token = request.headers.get("Authorization", "")
    token = token[7:].strip() if token.startswith("Bearer ") else (request.cookies.get("token") or request.cookies.get("jwt"))
    if not token:
        return jsonify({"success": False, "message": "Authentication required."}), 401
    try:
        payload = jwt.decode(token, current_app.config["SECRET_KEY"], algorithms=["HS256"])
        user = db.session.get(User, payload.get("user_id"))
    except Exception:
        user = None
    if not user or user.status != "active":
        return jsonify({"success": False, "message": "Authentication required."}), 401
    request.crop_user = user


def _error(message, status=400):
    return jsonify({"success": False, "message": message}), status


def _success(data, status=200, message=None):
    body = {"success": True, "data": data}
    if message:
        body["message"] = message
    return jsonify(body), status


def seed_crop_intelligence_references():
    """Seed default sourced profiles and planning rules without overwriting edits."""
    changed = False
    for code, profile in VARIETY_SEED.items():
        if not CropVarietyReference.query.filter_by(variety_code=code).first():
            db.session.add(CropVarietyReference(variety_code=code, profile_json=json_text(profile)))
            changed = True
    for stage in STAGE_SEED:
        if not CropStageReference.query.filter_by(stage_key=stage["key"]).first():
            db.session.add(CropStageReference(stage_key=stage["key"], stage_json=json_text(stage)))
            changed = True
    for key, strategy in MITIGATION_SEED.items():
        if not CropMitigationStrategy.query.filter_by(strategy_key=key).first():
            db.session.add(CropMitigationStrategy(strategy_key=key, strategy_json=json_text(strategy)))
            changed = True
    if changed:
        db.session.commit()


def _owned_cycle(user, cycle_id):
    return CropCycle.query.filter_by(id=cycle_id, user_id=user.id).first()


def _owned_field(user, field_id):
    return Field.query.join(Farm).filter(
        Field.id == field_id, Field.is_archived.is_(False),
        Farm.user_id == user.id, Farm.is_archived.is_(False),
    ).first()


def _json_references(model, key_column, value_column):
    rows = model.query.order_by(getattr(model, key_column)).all()
    return {getattr(row, key_column): parse_json(getattr(row, value_column)) for row in rows}


def _is_admin(user):
    return user.role == "admin"


def _require_admin():
    if not _is_admin(request.crop_user):
        return _error("Administrator access is required to edit crop reference data.", 403)
    return None


def _number(data, key, required=False):
    raw = data.get(key)
    if raw in (None, ""):
        if required:
            raise ValueError(f"{key.replace('_', ' ').title()} is required.")
        return None
    if isinstance(raw, bool):
        raise ValueError(f"{key.replace('_', ' ').title()} must be numeric.")
    try:
        value = float(raw)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{key.replace('_', ' ').title()} must be numeric.") from exc
    if not math.isfinite(value):
        raise ValueError(f"{key.replace('_', ' ').title()} must be finite.")
    low, high = NUMERIC_LIMITS.get(key, (-float("inf"), float("inf")))
    if value < low or value > high:
        raise ValueError(f"{key.replace('_', ' ').title()} must be between {low} and {high}.")
    if key == "measurement_area_m2" and value <= 0:
        raise ValueError("Measurement area must be greater than zero.")
    if key == "plant_height_cm" and value <= 0:
        raise ValueError("Plant height must be greater than zero.")
    if key == "tiller_count" and not value.is_integer():
        raise ValueError("Tiller count must be a whole number.")
    return int(value) if key == "tiller_count" else value


def _cycle_text(data, key, limit):
    value = data.get(key, "")
    if value is None:
        return None
    if not isinstance(value, str) or len(value) > limit:
        raise ValueError(f"{key.replace('_', ' ').title()} must be text under {limit} characters.")
    return value.strip() or None


def _validate_observation(data, cycle, source="manual"):
    if not isinstance(data, dict):
        raise ValueError("Observation must be a JSON object.")
    observation_date = parse_date(data.get("observation_date"), "Observation date")
    if observation_date < cycle.planting_date:
        raise ValueError("Observation date cannot be before the crop planting date.")
    if observation_date > date.today():
        raise ValueError("Observation date cannot be in the future.")
    result = {key: _number(data, key) for key in NUMERIC_LIMITS if key in data}
    for key in NUMERIC_LIMITS:
        if key not in result:
            result[key] = _number(data, key)
    if result["tiller_count"] is not None and result["measurement_area_m2"] is None:
        raise ValueError("Measurement area is required when recording a tiller count.")
    leaf = (data.get("leaf_condition") or "").strip().casefold() or None
    if leaf and leaf not in LEAF_CONDITIONS:
        raise ValueError("Select a supported leaf condition.")
    result.update({
        "observation_date": observation_date,
        "dap": dap_for(cycle.planting_date, observation_date),
        "leaf_condition": leaf,
        "source": source,
    })
    for key in OBSERVATION_KEYS - set(NUMERIC_LIMITS) - {"leaf_condition"}:
        value = data.get(key)
        if value is not None and not isinstance(value, str):
            raise ValueError(f"{key.replace('_', ' ').title()} must be text.")
        if isinstance(value, str) and len(value) > 4000:
            raise ValueError(f"{key.replace('_', ' ').title()} is too long.")
        result[key] = value.strip() if isinstance(value, str) else None
    return result


def _profile(code):
    row = CropVarietyReference.query.filter_by(variety_code=code).first()
    return parse_json(row.profile_json) if row else None


def _strategies():
    return _json_references(CropMitigationStrategy, "strategy_key", "strategy_json")


def _stages():
    return list(_json_references(CropStageReference, "stage_key", "stage_json").values())


def _alert_json(alert):
    payload = alert_dict(alert)
    observation = db.session.get(CropObservation, alert.observation_id) if alert.observation_id else None
    cycle = db.session.get(CropCycle, alert.cycle_id)
    if observation:
        payload.update({"observation_date": observation.observation_date.isoformat(), "dap": observation.dap})
    if cycle:
        farm, field = db.session.get(Farm, cycle.farm_id), db.session.get(Field, cycle.field_id)
        payload.update({"variety_code": cycle.variety_code, "farm_id": cycle.farm_id, "farm_name": farm.name if farm else None, "field_id": cycle.field_id, "field_name": field.name if field else None})
    return payload


def _sync_assessment(cycle, observation):
    profile = _profile(cycle.variety_code) or {}
    farm = db.session.get(Farm, cycle.farm_id)
    region = {"state": farm.state, "district": farm.district} if farm else None
    history = CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.asc()).all()
    assessment = assess_observation(cycle, observation, profile, history, _strategies(), region)
    for comparison in assessment["comparisons"]:
        row = CropComparison.query.filter_by(observation_id=observation.id, metric=comparison["metric"]).first()
        if row is None:
            row = CropComparison(observation_id=observation.id, metric=comparison["metric"])
            db.session.add(row)
        row.actual_value = float(comparison["actual"]) if comparison["actual"] is not None else None
        row.unit = comparison.get("unit")
        row.reference_status = comparison["status"]
        row.reference_json = json_text(comparison["reference"]) if comparison.get("reference") else None
        row.explanation = comparison["explanation"]
        row.updated_at = datetime.utcnow()
    active_rules = {item["rule_key"] for item in assessment["alerts"]}
    old_alerts = CropAlert.query.filter_by(cycle_id=cycle.id, observation_id=observation.id).all()
    for item in assessment["alerts"]:
        existing = next((alert for alert in old_alerts if alert.rule_key == item["rule_key"]), None)
        if existing:
            continue
        strategy = item
        db.session.add(CropAlert(
            cycle_id=cycle.id, observation_id=observation.id, rule_key=item["rule_key"],
            indicator=item["indicator"], observed_value=item["observed_value"], reference_value=item["reference_value"],
            severity=item["severity"], evidence=item["evidence"],
            possible_explanations_json=json_text(item["possible_explanations"]),
            investigate_json=json_text(item["investigate"]), mitigation_json=json_text({
                "actions": item["mitigation"], "avoid": item["avoid"], "expert_help": item["expert_help"],
                "follow_up_days": item["follow_up_days"], "source": item["source"],
            }),
            follow_up_date=observation.observation_date + timedelta(days=item["follow_up_days"]),
        ))
        db.session.add(CropRecommendation(
            cycle_id=cycle.id, observation_id=observation.id,
            recommendation_key=item["rule_key"],
            recommendation_json=json_text({
                "title": f"Investigate {item['indicator'].lower()}",
                "observed": item["observed_value"], "evidence": item["evidence"],
                "why": item["possible_explanations"], "check_first": item["investigate"],
                "actions": item["mitigation"], "avoid": item["avoid"],
                "expected_benefit": "Collecting comparable field evidence can help the farmer or local adviser choose an appropriate response.",
                "follow_up_days": item["follow_up_days"], "expert_help": item["expert_help"],
                "source": item["source"], "classification": "Rule-based scouting prompt; possible explanations only.",
            }),
        ))
    for alert in old_alerts:
        if alert.rule_key not in active_rules and alert.status not in {"resolved", "unconfirmed"}:
            old_status = alert.status
            alert.status = "unconfirmed"
            alert.resolution_notes = "The associated observation was corrected; this rule no longer matches the corrected record."
            alert.resolved_at = datetime.utcnow()
            db.session.add(CropAlertEvent(alert_id=alert.id, actor_id=request.crop_user.id, old_status=old_status, new_status="unconfirmed", note=alert.resolution_notes))
    return assessment


def _cycle_overview(cycle):
    farm = db.session.get(Farm, cycle.farm_id)
    field = db.session.get(Field, cycle.field_id)
    today = date.today()
    dap = dap_for(cycle.planting_date, today)
    stages = _stages()
    current = stage_for_dap(dap, stages)
    observations = CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.asc()).all()
    tasks = CropTask.query.filter_by(cycle_id=cycle.id).order_by(CropTask.due_date.asc()).all()
    last_observation = observations[-1] if observations else None
    assessment = _sync_assessment(cycle, last_observation) if last_observation else None
    if assessment:
        db.session.commit()
    alerts = CropAlert.query.filter_by(cycle_id=cycle.id).order_by(CropAlert.created_at.desc()).all()
    stage_summaries = [summarize_stage(cycle, stage, observations, tasks, alerts, today) for stage in stages]
    return {
        "cycle": {"id": cycle.id, "farm_id": cycle.farm_id, "farm_name": farm.name if farm else "", "farm_state": farm.state if farm else None, "farm_district": farm.district if farm else None, "field_id": cycle.field_id, "field_name": field.name if field else "", "variety_code": cycle.variety_code, "planting_date": cycle.planting_date.isoformat(), "planting_system": cycle.planting_system, "season": cycle.season, "notes": cycle.notes, "status": cycle.status, "created_at": cycle.created_at.isoformat() if cycle.created_at else None},
        "as_of_date": today.isoformat(),
        "actual_dap": dap,
        "current_stage": current,
        "stage_progress_pct": stage_progress(dap, current),
        "next_checkpoint": next_checkpoint(dap),
        "variety_reference": _profile(cycle.variety_code) or {},
        "stage_summaries": stage_summaries,
        "observations": [observation_dict(item) for item in observations],
        "tasks": [task_dict(item) for item in tasks],
        "alerts": [_alert_json(item) for item in alerts],
        "latest_assessment": assessment,
    }


@crop_intelligence_bp.get("/fields")
def list_monitorable_fields():
    user = request.crop_user
    fields = Field.query.join(Farm).filter(Field.is_archived.is_(False), Farm.is_archived.is_(False), Farm.user_id == user.id).order_by(Farm.name, Field.name).all()
    return _success({"fields": [{"id": field.id, "farm_id": field.farm_id, "farm_name": field.farm.name, "name": field.name, "state": field.farm.state, "district": field.farm.district, "variety_code": field.sugarcane_variety, "planting_date": field.planting_date.isoformat() if field.planting_date else None, "soil_type": field.soil_type, "soil_ph": float(field.soil_ph) if field.soil_ph is not None else None, "soil_moisture": float(field.soil_moisture) if field.soil_moisture is not None else None, "irrigation_method": field.irrigation_method} for field in fields]})


@crop_intelligence_bp.get("/varieties")
def list_crop_variety_references():
    # Start with stored profiles so custom fields (such as configured trial
    # ranges) survive, then overlay maintained source-backed fields from the
    # current seed. This upgrades stale seeded values without discarding extras.
    profiles = {}
    rows = CropVarietyReference.query.order_by(CropVarietyReference.variety_code).all()
    for row in rows:
        stored_profile = parse_json(row.profile_json)
        seed_profile = VARIETY_SEED.get(row.variety_code)
        profiles[row.variety_code] = {**stored_profile, **seed_profile} if seed_profile else stored_profile
    return _success({"varieties": [{"code": code, **profile} for code, profile in profiles.items()]})


@crop_intelligence_bp.put("/varieties/<path:variety_code>")
def update_crop_variety_reference(variety_code):
    denied = _require_admin()
    if denied:
        return denied
    row = CropVarietyReference.query.filter_by(variety_code=variety_code).first()
    data = request.get_json(silent=True)
    if not row or not isinstance(data, dict):
        return _error("Known variety and JSON profile are required.")
    ranges = data.get("growth_reference_ranges", [])
    if not isinstance(ranges, list):
        return _error("growth_reference_ranges must be a list.")
    for reference in ranges:
        try:
            low, high = float(reference["low"]), float(reference["high"])
            start, end = int(reference["start_day"]), int(reference["end_day"])
            source = reference.get("source")
            if (
                not all(math.isfinite(value) for value in (low, high)) or low > high
                or start < 0 or end < start or reference.get("metric") not in NUMERIC_METRICS
                or not reference.get("unit") or not reference.get("region")
                or not reference.get("conditions") or not reference.get("method")
                or not isinstance(source, dict) or not source.get("title")
                or not isinstance(source.get("url"), str) or not source["url"].startswith(("https://", "http://"))
                or not source.get("verified")
            ):
                raise ValueError
            if reference.get("metric") == "tiller_count" and reference.get("unit") not in {"tillers/m²", "tillers per m²"}:
                raise ValueError
        except (KeyError, TypeError, ValueError):
            return _error("Each reference range needs finite bounds, crop-age limits, unit, exact region, conditions, measurement method, and a titled source with URL and verification date.")
    row.profile_json = json_text(data)
    row.updated_by = request.crop_user.id
    row.updated_at = datetime.utcnow()
    db.session.commit()
    return _success({"code": variety_code, **data}, message="Variety reference saved.")


@crop_intelligence_bp.get("/stages")
def list_growth_stages():
    return _success({"stages": _stages()})


@crop_intelligence_bp.put("/stages/<stage_key>")
def update_growth_stage(stage_key):
    denied = _require_admin()
    if denied:
        return denied
    row = CropStageReference.query.filter_by(stage_key=stage_key).first()
    data = request.get_json(silent=True)
    if not row or not isinstance(data, dict) or data.get("key") != stage_key:
        return _error("Known stage key and matching JSON configuration are required.")
    try:
        start = int(data["start_day"])
        end = int(data["end_day"]) if data.get("end_day") is not None else None
        if start < 1 or (end is not None and end < start) or not data.get("name"):
            raise ValueError
    except (TypeError, KeyError, ValueError):
        return _error("Stage needs a name, start day >= 1, and an end day at or after its start, or null for continued monitoring.")
    row.stage_json = json_text(data)
    row.updated_by = request.crop_user.id
    row.updated_at = datetime.utcnow()
    db.session.commit()
    return _success(data, message="Growth stage reference saved.")


@crop_intelligence_bp.get("/mitigations")
def list_mitigation_references():
    return _success({"strategies": _json_references(CropMitigationStrategy, "strategy_key", "strategy_json")})


@crop_intelligence_bp.put("/mitigations/<strategy_key>")
def update_mitigation_reference(strategy_key):
    denied = _require_admin()
    if denied:
        return denied
    row = CropMitigationStrategy.query.filter_by(strategy_key=strategy_key).first()
    data = request.get_json(silent=True)
    if not row or not isinstance(data, dict) or not data.get("problem"):
        return _error("Known strategy key and problem description are required.")
    for name in ("investigate", "immediate_actions", "avoid"):
        if not isinstance(data.get(name), list) or not all(isinstance(value, str) for value in data[name]):
            return _error(f"{name} must be a list of text steps.")
    row.strategy_json = json_text(data)
    row.updated_by = request.crop_user.id
    row.updated_at = datetime.utcnow()
    db.session.commit()
    return _success(data, message="Mitigation strategy saved.")


@crop_intelligence_bp.route("/cycles", methods=["GET", "POST"])
def crop_cycles():
    user = request.crop_user
    if request.method == "GET":
        cycles = CropCycle.query.filter_by(user_id=user.id).order_by(CropCycle.status.asc(), CropCycle.updated_at.desc()).all()
        return _success({"cycles": [_cycle_overview(cycle) for cycle in cycles]})
    data = request.get_json(silent=True) or {}
    try:
        field_id = int(data.get("field_id"))
        field = _owned_field(user, field_id)
        if not field:
            return _error("Owned, active field not found.", 404)
        planting_date = parse_date(data.get("planting_date") or (field.planting_date.isoformat() if field.planting_date else None), "Planting date")
        if planting_date > date.today():
            return _error("Planting date cannot be in the future.")
        code = (data.get("variety_code") or field.sugarcane_variety or "").strip()
        if not _profile(code):
            return _error("Select a sugarcane variety supported by crop monitoring.")
        cycle = CropCycle(
            user_id=user.id, farm_id=field.farm_id, field_id=field.id, variety_code=code,
            planting_date=planting_date,
            planting_system=_cycle_text(data, "planting_system", 80),
            season=_cycle_text(data, "season", 80),
            notes=_cycle_text(data, "notes", 2000),
        )
        db.session.add(cycle)
        db.session.commit()
        return _success(_cycle_overview(cycle), 201, "Crop cycle created.")
    except (ValueError, TypeError) as exc:
        db.session.rollback()
        return _error(str(exc))


@crop_intelligence_bp.route("/cycles/<int:cycle_id>", methods=["GET", "PATCH"])
def crop_cycle(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    if request.method == "GET":
        return _success(_cycle_overview(cycle))
    data = request.get_json(silent=True) or {}
    try:
        if "status" in data and data["status"] not in {"active", "archived"}:
            return _error("Cycle status must be active or archived.")
        for key, limit in (("planting_system", 80), ("season", 80), ("notes", 2000)):
            if key in data:
                _cycle_text(data, key, limit)
        if "planting_date" in data:
            planting_date = parse_date(data["planting_date"], "Planting date")
            if planting_date > date.today():
                return _error("Planting date cannot be in the future.")
            observations = CropObservation.query.filter_by(cycle_id=cycle.id).all()
            if any(item.observation_date < planting_date for item in observations):
                return _error("Planting date cannot move after an existing observation.")
            cycle.planting_date = planting_date
            for item in observations:
                item.dap = dap_for(planting_date, item.observation_date)
        if "variety_code" in data:
            code = str(data["variety_code"]).strip()
            if not _profile(code):
                return _error("Select a supported sugarcane variety.")
            cycle.variety_code = code
        for key in ("planting_system", "season", "notes"):
            if key in data:
                setattr(cycle, key, _cycle_text(data, key, 2000 if key == "notes" else 80))
        if "status" in data:
            cycle.status = data["status"]
        cycle.updated_at = datetime.utcnow()
        db.session.commit()
    except ValueError as exc:
        db.session.rollback()
        return _error(str(exc))
    return _success(_cycle_overview(cycle), message="Crop cycle updated; DAP values were recalculated.")


@crop_intelligence_bp.delete("/cycles/<int:cycle_id>")
def archive_crop_cycle(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    cycle.status = "archived"
    cycle.updated_at = datetime.utcnow()
    db.session.commit()
    return _success(_cycle_overview(cycle), message="Crop cycle archived. Its observations and history are preserved.")


@crop_intelligence_bp.route("/cycles/<int:cycle_id>/observations", methods=["GET", "POST"])
def crop_observations(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    if request.method == "GET":
        query = CropObservation.query.filter_by(cycle_id=cycle.id)
        start, end = request.args.get("from"), request.args.get("to")
        try:
            if start:
                query = query.filter(CropObservation.observation_date >= parse_date(start, "from date"))
            if end:
                query = query.filter(CropObservation.observation_date <= parse_date(end, "to date"))
        except ValueError as exc:
            return _error(str(exc))
        return _success({"observations": [observation_dict(item) for item in query.order_by(CropObservation.observation_date.asc()).limit(1000).all()]})
    try:
        normalized = _validate_observation(request.get_json(silent=True) or {}, cycle)
        existing = CropObservation.query.filter_by(cycle_id=cycle.id, observation_date=normalized["observation_date"]).first()
        if existing:
            return _error("An observation already exists for this date. Use PATCH to correct it.", 409)
        observation = CropObservation(cycle_id=cycle.id, created_by=request.crop_user.id, **normalized)
        db.session.add(observation)
        db.session.flush()
        assessment = _sync_assessment(cycle, observation)
        db.session.commit()
        return _success({"observation": observation_dict(observation), "assessment": assessment}, 201, "Observation saved and assessed.")
    except ValueError as exc:
        db.session.rollback()
        return _error(str(exc))


@crop_intelligence_bp.patch("/cycles/<int:cycle_id>/observations/<int:observation_id>")
def correct_observation(cycle_id, observation_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    observation = CropObservation.query.filter_by(id=observation_id, cycle_id=cycle.id).first() if cycle else None
    if not cycle or not observation:
        return _error("Observation not found.", 404)
    data = request.get_json(silent=True) or {}
    allowed = OBSERVATION_KEYS | {"observation_date"}
    if set(data) - allowed:
        return _error("Only observation fields may be corrected.")
    payload = {key: getattr(observation, key) for key in OBSERVATION_FIELDS}
    payload["observation_date"] = observation.observation_date.isoformat()
    payload.update(data)
    try:
        normalized = _validate_observation(payload, cycle, source=observation.source)
        conflict = CropObservation.query.filter_by(cycle_id=cycle.id, observation_date=normalized["observation_date"]).filter(CropObservation.id != observation.id).first()
        if conflict:
            return _error("Another observation already uses that date.", 409)
        old_snapshot = {key: getattr(observation, key) for key in OBSERVATION_FIELDS}
        old_snapshot.update({"observation_date": observation.observation_date.isoformat(), "dap": observation.dap})
        db.session.add(CropObservationRevision(observation_id=observation.id, changed_by=request.crop_user.id, snapshot_json=json_text(old_snapshot)))
        for key, value in normalized.items():
            setattr(observation, key, value)
        observation.updated_at = datetime.utcnow()
        assessment = _sync_assessment(cycle, observation)
        db.session.commit()
        return _success({"observation": observation_dict(observation), "assessment": assessment}, message="Observation corrected; prior values are retained in revision history.")
    except ValueError as exc:
        db.session.rollback()
        return _error(str(exc))


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/observations/<int:observation_id>/revisions")
def observation_revisions(cycle_id, observation_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    observation = CropObservation.query.filter_by(id=observation_id, cycle_id=cycle.id).first() if cycle else None
    if not observation:
        return _error("Observation not found.", 404)
    revisions = CropObservationRevision.query.filter_by(observation_id=observation.id).order_by(CropObservationRevision.created_at.desc()).all()
    return _success({"revisions": [{"id": item.id, "changed_by": item.changed_by, "snapshot": parse_json(item.snapshot_json), "created_at": item.created_at.isoformat()} for item in revisions]})


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/calendar")
def crop_calendar(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    try:
        anchor = parse_date(request.args.get("date") or date.today().isoformat(), "Calendar date")
        view = request.args.get("view", "week")
        if view not in {"day", "week", "stage"}:
            return _error("Calendar view must be day, week, or stage.")
        if request.args.get("from") or request.args.get("to"):
            start = parse_date(request.args.get("from") or anchor.isoformat(), "from date")
            end = parse_date(request.args.get("to") or anchor.isoformat(), "to date")
        elif view == "day":
            start = end = anchor
        elif view == "stage":
            dap = dap_for(cycle.planting_date, anchor)
            stage = stage_for_dap(dap, _stages())
            if stage is None:
                start = end = anchor
            else:
                start = cycle.planting_date + timedelta(days=int(stage["start_day"]))
                end = cycle.planting_date + timedelta(days=int(stage["end_day"])) if stage.get("end_day") is not None else start + timedelta(days=29)
        else:
            start = anchor - timedelta(days=anchor.weekday())
            end = start + timedelta(days=6)
        if end < start or (end - start).days > 61:
            return _error("Calendar date range must be 1–62 days.")
    except (TypeError, ValueError) as exc:
        return _error(str(exc))
    stages = _stages()
    profile = _profile(cycle.variety_code) or {}
    observations = {item.observation_date: item for item in CropObservation.query.filter(CropObservation.cycle_id == cycle.id, CropObservation.observation_date >= start, CropObservation.observation_date <= end).all()}
    tasks = CropTask.query.filter(CropTask.cycle_id == cycle.id, CropTask.due_date >= start, CropTask.due_date <= end).order_by(CropTask.due_date).all()
    task_dates = {}
    for task in tasks:
        task_dates.setdefault(task.due_date, []).append(task_dict(task))
    activities = CropActivity.query.filter(CropActivity.cycle_id == cycle.id, CropActivity.activity_date >= start, CropActivity.activity_date <= end).order_by(CropActivity.activity_date).all()
    activity_dates = {}
    for activity in activities:
        activity_dates.setdefault(activity.activity_date, []).append({"id": activity.id, "activity_date": activity.activity_date.isoformat(), "activity_type": activity.activity_type, "details": activity.details, "amount": activity.amount, "amount_unit": activity.amount_unit})
    alerts = CropAlert.query.filter_by(cycle_id=cycle.id).filter(CropAlert.status.notin_(["resolved", "unconfirmed"])).all()
    alert_dates = {}
    for alert in alerts:
        if alert.observation_id:
            obs = db.session.get(CropObservation, alert.observation_id)
            if obs:
                alert_dates.setdefault(obs.observation_date, []).append(_alert_json(alert))
    days = []
    cursor = start
    while cursor <= end:
        dap, stage = stage_for_calendar_day(cycle, cursor, stages)
        observation = observations.get(cursor)
        expected_observation = dap > 0 and dap % 7 == 0
        days.append({
            "date": cursor.isoformat(), "dap": dap,
            "stage": stage,
            "developmental_activity": stage.get("process") if stage else "Before planting date",
            "recommended_observation": stage.get("observations", [None])[0] if stage else None,
            "observation_frequency": stage.get("frequency") if stage else None,
            "observation": observation_dict(observation) if observation else None,
            "tasks": task_dates.get(cursor, []), "activities": activity_dates.get(cursor, []), "alerts": alert_dates.get(cursor, []),
            "is_recorded": observation is not None,
            "missed_observation": bool(cursor < date.today() and expected_observation and not observation),
            "stage_guidance": expected_stage_guidance(stage, profile) if stage else None,
        })
        cursor += timedelta(days=1)
    return _success({"cycle_id": cycle.id, "view": view, "from": start.isoformat(), "to": end.isoformat(), "days": days})


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/comparison")
def crop_growth_comparison(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    observations = CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.asc()).all()
    farm = db.session.get(Farm, cycle.farm_id)
    region = {"state": farm.state, "district": farm.district} if farm else None
    comparisons = []
    for observation in observations:
        assessment = assess_observation(cycle, observation, _profile(cycle.variety_code) or {}, observations, _strategies(), region)
        comparisons.append({"observation_id": observation.id, "observation_date": observation.observation_date.isoformat(), "dap": observation.dap, "source": observation.source, "measurements": assessment["comparisons"], "data_quality": assessment["data_quality"]})
    return _success({"variety_code": cycle.variety_code, "reference_status": "Only sourced ranges matching this variety and crop age are compared. Unavailable ranges stay unscored.", "comparisons": comparisons})


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/recommendations")
def crop_recommendations(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    farm, field = db.session.get(Farm, cycle.farm_id), db.session.get(Field, cycle.field_id)
    dap = dap_for(cycle.planting_date)
    stage = stage_for_dap(dap, _stages())
    last = CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.desc()).first()
    active_alerts = CropAlert.query.filter_by(cycle_id=cycle.id).filter(CropAlert.status.notin_(["resolved", "unconfirmed"])).all()
    profile = _profile(cycle.variety_code) or {}
    activities = []
    for item in (stage or {}).get("activities", []):
        activities.append({
            **item,
            "reason": item.get("why"),
            "evidence": {
                "crop_age_dap": dap, "growth_stage": (stage or {}).get("name"),
                "variety": cycle.variety_code, "state": farm.state if farm else None,
                "district": farm.district if farm else None,
                "soil_type": field.soil_type if field else None,
                "soil_ph": float(field.soil_ph) if field and field.soil_ph is not None else None,
                "irrigation_method": field.irrigation_method if field else None,
                "latest_observation_date": last.observation_date.isoformat() if last else None,
                "latest_soil_moisture_pct": last.soil_moisture_pct if last else None,
                "latest_rainfall_mm": last.rainfall_mm if last else None,
            },
            "classification": "General stage-planning guidance",
            "source": {"organization": "ICAR / local agricultural extension", "title": "ICAR Kharif Agro-Advisories for Farmers 2025", "url": "https://icar.gov.in/en/icar-kharif-agro-advisories-farmers-2025-1", "limitation": "General guidance; confirm local season, soil, and variety recommendations."},
        })
    stored = CropRecommendation.query.filter_by(cycle_id=cycle.id).order_by(CropRecommendation.created_at.desc()).limit(100).all()
    return _success({
        "cycle_id": cycle.id, "stage": stage, "variety_reference": profile,
        "recommendations": activities,
        "evidence_based_follow_ups": [{"id": row.id, **parse_json(row.recommendation_json), "created_at": row.created_at.isoformat()} for row in stored],
        "active_alerts": [_alert_json(item) for item in active_alerts],
        "limitations": ["No forecast was available to this module when the assessment was made.", "No fertilizer or pesticide dose is prescribed.", "A field observation is not a confirmed diagnosis."],
    })


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/tasks")
def list_crop_tasks(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    rows = CropTask.query.filter_by(cycle_id=cycle.id).order_by(CropTask.due_date.asc(), CropTask.id.asc()).limit(1000).all()
    return _success({"tasks": [task_dict(item) for item in rows]})


@crop_intelligence_bp.post("/cycles/<int:cycle_id>/tasks")
def create_crop_task(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    data = request.get_json(silent=True) or {}
    try:
        due = parse_date(data.get("due_date"), "Due date")
    except ValueError as exc:
        return _error(str(exc))
    title = data.get("title")
    if not isinstance(title, str) or not title.strip() or len(title) > 180:
        return _error("Task title must contain 1–180 characters.")
    category = data.get("category", "scouting")
    priority = data.get("priority", "normal")
    details = data.get("details", "")
    if category not in {"irrigation", "nutrition", "weeds", "pest_disease", "soil", "scouting", "assessment", "other"}:
        return _error("Select a supported task category.")
    if priority not in {"low", "normal", "high"} or not isinstance(details, str) or len(details) > 3000:
        return _error("Task priority or details are invalid.")
    key = hashlib.sha256(f"{due.isoformat()}|{title.strip().casefold()}".encode()).hexdigest()[:64]
    existing = CropTask.query.filter_by(cycle_id=cycle.id, idempotency_key=key).first()
    if existing:
        return _success({"task": task_dict(existing), "duplicate": True}, message="This task already exists.")
    task = CropTask(cycle_id=cycle.id, idempotency_key=key, due_date=due, title=title.strip(), category=category, details=details.strip() or None, priority=priority, status="pending", created_by=request.crop_user.id)
    db.session.add(task)
    db.session.flush()
    db.session.add(CropTaskEvent(task_id=task.id, actor_id=request.crop_user.id, old_status=None, new_status="pending", note="Task created."))
    db.session.commit()
    return _success({"task": task_dict(task), "duplicate": False}, 201, "Task scheduled.")


@crop_intelligence_bp.patch("/cycles/<int:cycle_id>/tasks/<int:task_id>")
def update_crop_task(cycle_id, task_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    task = CropTask.query.filter_by(id=task_id, cycle_id=cycle.id).first() if cycle else None
    if not task:
        return _error("Task not found.", 404)
    data = request.get_json(silent=True) or {}
    old_status = task.status
    if "status" in data:
        if data["status"] not in TASK_STATUSES:
            return _error("Invalid task status.")
        task.status = data["status"]
        task.completed_at = datetime.utcnow() if task.status == "completed" else None
    if "details" in data:
        if data["details"] is not None and (not isinstance(data["details"], str) or len(data["details"]) > 3000):
            return _error("Task details must be text under 3000 characters.")
        task.details = data["details"]
    if "due_date" in data:
        try:
            task.due_date = parse_date(data["due_date"], "Due date")
        except ValueError as exc:
            return _error(str(exc))
    task.updated_at = datetime.utcnow()
    if old_status != task.status:
        db.session.add(CropTaskEvent(task_id=task.id, actor_id=request.crop_user.id, old_status=old_status, new_status=task.status, note=data.get("note")))
    db.session.commit()
    return _success({"task": task_dict(task)})


@crop_intelligence_bp.route("/cycles/<int:cycle_id>/activities", methods=["GET", "POST"])
def crop_activities(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    if request.method == "GET":
        rows = CropActivity.query.filter_by(cycle_id=cycle.id).order_by(CropActivity.activity_date.desc()).limit(1000).all()
        return _success({"activities": [{"id": row.id, "activity_date": row.activity_date.isoformat(), "activity_type": row.activity_type, "details": row.details, "amount": row.amount, "amount_unit": row.amount_unit, "created_at": row.created_at.isoformat()} for row in rows]})
    data = request.get_json(silent=True) or {}
    try:
        activity_date = parse_date(data.get("activity_date"), "Activity date")
        if activity_date < cycle.planting_date or activity_date > date.today():
            return _error("Activity date must be on or after planting and no later than today.")
        kind = data.get("activity_type")
        details = data.get("details")
        if kind not in {"irrigation", "nutrition", "weeding", "scouting", "soil", "pest_disease", "other"} or not isinstance(details, str) or not details.strip() or len(details) > 3000:
            return _error("Select an activity type and provide details under 3000 characters.")
        amount = _number(data, "amount")
        amount_unit = data.get("amount_unit")
        if amount_unit is not None and (not isinstance(amount_unit, str) or len(amount_unit) > 32):
            return _error("Activity unit must be text under 32 characters.")
    except ValueError as exc:
        return _error(str(exc))
    activity = CropActivity(cycle_id=cycle.id, activity_date=activity_date, activity_type=kind, details=details.strip(), amount=amount, amount_unit=amount_unit, created_by=request.crop_user.id)
    db.session.add(activity)
    db.session.commit()
    return _success({"activity": {"id": activity.id, "activity_date": activity.activity_date.isoformat(), "activity_type": activity.activity_type, "details": activity.details, "amount": activity.amount, "amount_unit": activity.amount_unit}}, 201)


@crop_intelligence_bp.route("/cycles/<int:cycle_id>/alerts", methods=["GET", "POST"])
def crop_alerts(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    if request.method == "GET":
        rows = CropAlert.query.filter_by(cycle_id=cycle.id).order_by(CropAlert.created_at.desc()).limit(1000).all()
        return _success({"alerts": [_alert_json(item) for item in rows]})
    latest = CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.desc()).all()
    created_before = CropAlert.query.filter_by(cycle_id=cycle.id).count()
    for observation in latest:
        _sync_assessment(cycle, observation)
    db.session.commit()
    rows = CropAlert.query.filter_by(cycle_id=cycle.id).order_by(CropAlert.created_at.desc()).all()
    return _success({"alerts": [_alert_json(item) for item in rows], "new_alerts": max(0, len(rows) - created_before)})


@crop_intelligence_bp.patch("/cycles/<int:cycle_id>/alerts/<int:alert_id>")
def update_crop_alert(cycle_id, alert_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    alert = CropAlert.query.filter_by(id=alert_id, cycle_id=cycle.id).first() if cycle else None
    if not alert:
        return _error("Alert not found.", 404)
    data = request.get_json(silent=True) or {}
    new_status = data.get("status")
    if new_status not in ALERT_STATUSES:
        return _error("Select a supported alert status.")
    notes = data.get("resolution_notes", data.get("note", ""))
    if not isinstance(notes, str) or len(notes) > 3000:
        return _error("Resolution notes must be text under 3000 characters.")
    old_status = alert.status
    alert.status = new_status
    alert.resolution_notes = notes.strip() or alert.resolution_notes
    alert.resolved_at = datetime.utcnow() if new_status in {"resolved", "unconfirmed"} else None
    db.session.add(CropAlertEvent(alert_id=alert.id, actor_id=request.crop_user.id, old_status=old_status, new_status=new_status, note=notes.strip() or None))
    db.session.commit()
    return _success({"alert": _alert_json(alert)})


def _read_import(file_storage):
    filename = secure_filename(file_storage.filename or "")
    ext = Path(filename).suffix.casefold()
    raw = file_storage.read(MAX_UPLOAD_BYTES + 1)
    if not filename or len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("Choose a file smaller than 5 MB.")
    if ext == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV must use UTF-8 encoding.") from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        rows = list(reader)
    elif ext == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 200 or sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError("Excel workbook expands beyond the 25 MB processing limit.")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(values, ())]
            rows = [dict(zip(headers, row)) for row in values]
            workbook.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Excel file could not be read. Upload a valid .xlsx workbook.") from exc
    else:
        raise ValueError("Upload a .csv or .xlsx file.")
    if len(headers) > 40 or len(rows) > MAX_IMPORT_ROWS:
        raise ValueError("Import files may contain up to 40 columns and 5,000 data rows.")
    rows = [row for row in rows if any(value not in (None, "") for value in row.values())]
    if not headers or not rows:
        raise ValueError("The file must contain a header and at least one non-empty data row.")
    return filename, raw, headers, rows


IMPORT_ALIASES = {
    "observation_date": {"observation_date", "date", "recorded_date"},
    "days_after_planting": {"days_after_planting", "dap", "crop_age_days"},
    "plant_height_cm": {"plant_height_cm", "height_cm", "plant_height"},
    "tiller_count": {"tiller_count", "tillers", "shoot_count"},
    "measurement_area_m2": {"measurement_area_m2", "area_m2", "sample_area_m2"},
    "establishment_pct": {"establishment_pct", "germination_pct", "establishment_percent"},
    "leaf_condition": {"leaf_condition", "leaf_status"},
    "leaf_symptoms": {"leaf_symptoms", "symptoms"},
    "soil_moisture_pct": {"soil_moisture_pct", "soil_moisture_percent"},
    "rainfall_mm": {"rainfall_mm", "rain_mm"},
    "irrigation_activity": {"irrigation_activity", "irrigation"},
    "pest_symptoms": {"pest_symptoms", "pest"},
    "disease_symptoms": {"disease_symptoms", "disease"},
    "soil_or_drainage": {"soil_or_drainage", "drainage"},
    "notes": {"notes", "farmer_notes", "comments"},
}


def _import_mapping(headers, requested):
    normalized = {re.sub(r"[^a-z0-9]+", "_", header.casefold()).strip("_"): header for header in headers}
    mapping = {}
    for canonical, aliases in IMPORT_ALIASES.items():
        requested_header = requested.get(canonical) if isinstance(requested, dict) else None
        if requested_header:
            if requested_header not in headers:
                raise ValueError(f"Mapped column for {canonical} was not found in the uploaded file.")
            mapping[canonical] = requested_header
            continue
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match:
            mapping[canonical] = match
    if "observation_date" not in mapping:
        raise ValueError("Map a column to observation_date.")
    return mapping


def _normalize_import_row(raw_row, row_number, mapping, cycle):
    raw = {key: raw_row.get(header) for key, header in mapping.items()}
    try:
        date_value = raw.get("observation_date")
        if isinstance(date_value, datetime):
            date_value = date_value.date().isoformat()
        elif isinstance(date_value, date):
            date_value = date_value.isoformat()
        payload = {key: value for key, value in raw.items() if key not in {"observation_date", "days_after_planting"}}
        payload["observation_date"] = str(date_value).strip() if date_value is not None else None
        normalized = _validate_observation(payload, cycle, source="imported")
        supplied_dap = raw.get("days_after_planting")
        if supplied_dap not in (None, "") and int(float(supplied_dap)) != normalized["dap"]:
            raise ValueError(f"Supplied DAP {supplied_dap} does not match calculated DAP {normalized['dap']}.")
        return normalized, None
    except (ValueError, TypeError) as exc:
        return None, f"Row {row_number}: {exc}"


def _parse_import_request():
    file_storage = request.files.get("file")
    if not file_storage:
        raise ValueError("Choose a CSV or Excel file.")
    cycle_id = request.form.get("cycle_id", type=int)
    cycle = _owned_cycle(request.crop_user, cycle_id) if cycle_id else None
    if not cycle:
        raise ValueError("Select one of your crop cycles.")
    filename, raw, headers, rows = _read_import(file_storage)
    try:
        mapping = json.loads(request.form.get("mapping", "{}"))
    except ValueError as exc:
        raise ValueError("Column mapping must be valid JSON.") from exc
    mapping = _import_mapping(headers, mapping)
    normalized = []
    errors = []
    seen_dates = set()
    for index, raw_row in enumerate(rows, start=2):
        row, error = _normalize_import_row(raw_row, index, mapping, cycle)
        if error:
            errors.append(error)
            continue
        if row["observation_date"] in seen_dates:
            errors.append(f"Row {index}: duplicate observation date in this file.")
            continue
        seen_dates.add(row["observation_date"])
        normalized.append(row)
    return cycle, filename, raw, headers, mapping, normalized, errors


@crop_intelligence_bp.post("/imports/preview")
def preview_crop_import():
    try:
        cycle, filename, raw, headers, mapping, rows, errors = _parse_import_request()
    except ValueError as exc:
        return _error(str(exc))
    duplicates = {item.observation_date.isoformat() for item in CropObservation.query.filter_by(cycle_id=cycle.id).all()}
    for row in rows:
        if row["observation_date"].isoformat() in duplicates:
            errors.append(f"{row['observation_date'].isoformat()}: an observation already exists for this date.")
    return _success({"filename": filename, "file_sha256": hashlib.sha256(raw).hexdigest(), "headers": headers, "mapping": mapping, "preview": [{**{key: (value.isoformat() if isinstance(value, date) else value) for key, value in row.items() if key in OBSERVATION_KEYS or key in {"observation_date", "dap", "source"}}} for row in rows[:20]], "valid_rows": max(0, len(rows) - sum(1 for row in rows if row["observation_date"].isoformat() in duplicates)), "rejected_rows": len(errors), "errors": errors[:100], "already_imported": CropImport.query.filter_by(cycle_id=cycle.id, file_sha256=hashlib.sha256(raw).hexdigest()).first() is not None})


@crop_intelligence_bp.post("/imports/confirm")
def confirm_crop_import():
    try:
        cycle, filename, raw, headers, mapping, rows, errors = _parse_import_request()
    except ValueError as exc:
        return _error(str(exc))
    digest = hashlib.sha256(raw).hexdigest()
    previous = CropImport.query.filter_by(cycle_id=cycle.id, file_sha256=digest).first()
    if previous:
        return _success({"import_id": previous.id, "imported_rows": previous.imported_rows, "rejected_rows": previous.rejected_rows, "duplicate_import": True, "errors": parse_json(previous.summary_json, {}).get("errors", [])}, message="This file was already imported; no duplicate records were created.")
    existing_dates = {item.observation_date for item in CropObservation.query.filter_by(cycle_id=cycle.id).all()}
    accepted = []
    for row in rows:
        if row["observation_date"] in existing_dates:
            errors.append(f"{row['observation_date'].isoformat()}: an observation already exists for this date.")
        else:
            existing_dates.add(row["observation_date"])
            accepted.append(row)
    if not accepted:
        return _error("No valid, non-duplicate rows are available to import.", 409)
    try:
        imported = []
        for payload in accepted:
            observation = CropObservation(cycle_id=cycle.id, created_by=request.crop_user.id, **payload)
            db.session.add(observation)
            imported.append(observation)
        db.session.flush()
        for observation in imported:
            _sync_assessment(cycle, observation)
        summary = {"filename": filename, "errors": errors[:200], "headers": headers, "mapping": mapping}
        record = CropImport(cycle_id=cycle.id, file_sha256=digest, filename=filename[:255], imported_rows=len(imported), rejected_rows=len(errors), summary_json=json_text(summary), created_by=request.crop_user.id)
        db.session.add(record)
        db.session.commit()
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Crop data import failed for cycle %s", cycle.id)
        return _error("Import could not be saved. No records were committed.", 500)
    return _success({"import_id": record.id, "imported_rows": record.imported_rows, "rejected_rows": record.rejected_rows, "duplicate_import": False, "errors": errors[:200]}, 201, "Validated crop data imported.")


@crop_intelligence_bp.get("/imports/template.csv")
def crop_import_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["observation_date", "days_after_planting", "plant_height_cm", "tiller_count", "measurement_area_m2", "establishment_pct", "leaf_condition", "leaf_symptoms", "soil_moisture_pct", "rainfall_mm", "irrigation_activity", "pest_symptoms", "disease_symptoms", "soil_or_drainage", "notes"])
    return send_file(io.BytesIO(output.getvalue().encode("utf-8-sig")), mimetype="text/csv", as_attachment=True, download_name="sugaryield_crop_observations_template.csv")


@crop_intelligence_bp.post("/cycles/<int:cycle_id>/observations/<int:observation_id>/photo")
def upload_observation_photo(cycle_id, observation_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    observation = CropObservation.query.filter_by(id=observation_id, cycle_id=cycle.id).first() if cycle else None
    photo = request.files.get("photo")
    if not observation or not photo:
        return _error("Observation and photo are required.")
    raw = photo.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        return _error("Photo must be smaller than 5 MB.")
    try:
        image = Image.open(io.BytesIO(raw))
        if image.width * image.height > 40_000_000:
            return _error("Photo dimensions exceed the supported 40-megapixel limit.")
        image.verify()
        extension = {"JPEG": ".jpg", "PNG": ".png", "WEBP": ".webp"}.get(image.format)
        if not extension:
            raise ValueError
    except Exception:
        return _error("Photo must be a valid JPEG, PNG, or WebP image.")
    storage_dir = Path(current_app.instance_path) / "crop_uploads"
    storage_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{uuid.uuid4().hex}{extension}"
    (storage_dir / filename).write_bytes(raw)
    observation.photo_path = filename
    observation.updated_at = datetime.utcnow()
    db.session.commit()
    return _success({"observation_id": observation.id, "photo_available": True}, 201)


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/observations/<int:observation_id>/photo")
def get_observation_photo(cycle_id, observation_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    observation = CropObservation.query.filter_by(id=observation_id, cycle_id=cycle.id).first() if cycle else None
    if not observation or not observation.photo_path or Path(observation.photo_path).name != observation.photo_path:
        return _error("Photo not found.", 404)
    return send_from_directory(Path(current_app.instance_path) / "crop_uploads", observation.photo_path, as_attachment=False)


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/reports")
def list_crop_reports(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    rows = CropProgressReport.query.filter_by(cycle_id=cycle.id).order_by(CropProgressReport.checkpoint_day.asc()).all()
    return _success({"reports": [{"id": row.id, "checkpoint_day": row.checkpoint_day, "generated_at": row.generated_at.isoformat(), "report": parse_json(row.report_json)} for row in rows]})


@crop_intelligence_bp.post("/cycles/<int:cycle_id>/reports")
def generate_crop_report(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    data = request.get_json(silent=True) or {}
    try:
        checkpoint = int(data.get("checkpoint_day"))
    except (TypeError, ValueError):
        return _error("Select a checkpoint day: 30, 60, 90, 120, 150, or 180.")
    if checkpoint not in CHECKPOINTS:
        return _error("Select a checkpoint day: 30, 60, 90, 120, 150, or 180.")
    if dap_for(cycle.planting_date) < checkpoint:
        return _error("A checkpoint report is available after the crop reaches that day.", 409)
    existing = CropProgressReport.query.filter_by(cycle_id=cycle.id, checkpoint_day=checkpoint).first()
    farm, field = db.session.get(Farm, cycle.farm_id), db.session.get(Field, cycle.field_id)
    observations = CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.asc()).all()
    tasks = CropTask.query.filter_by(cycle_id=cycle.id).order_by(CropTask.due_date.asc()).all()
    alerts = CropAlert.query.filter_by(cycle_id=cycle.id).order_by(CropAlert.created_at.asc()).all()
    report = report_payload(cycle, farm, field, _stages(), observations, tasks, alerts, _profile(cycle.variety_code) or {}, _strategies())
    report["alerts"] = [_alert_json(item) for item in alerts]
    report["checkpoint_day"] = checkpoint
    report["checkpoint_assessment"] = "Evidence report generated from persisted farmer observations."
    if existing:
        existing.report_json = json_text(report)
        existing.generated_at = datetime.utcnow()
        existing.generated_by = request.crop_user.id
        row = existing
    else:
        row = CropProgressReport(cycle_id=cycle.id, checkpoint_day=checkpoint, report_json=json_text(report), generated_by=request.crop_user.id)
        db.session.add(row)
    db.session.commit()
    return _success({"id": row.id, "checkpoint_day": row.checkpoint_day, "generated_at": row.generated_at.isoformat(), "report": report}, 201 if not existing else 200, "Progress report saved.")


@crop_intelligence_bp.get("/cycles/<int:cycle_id>/export.csv")
def export_crop_history(cycle_id):
    cycle = _owned_cycle(request.crop_user, cycle_id)
    if not cycle:
        return _error("Crop cycle not found.", 404)
    output = io.StringIO()
    headers = ["observation_date", "dap", *OBSERVATION_FIELDS, "source"]
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for observation in CropObservation.query.filter_by(cycle_id=cycle.id).order_by(CropObservation.observation_date.asc()).all():
        row = observation_dict(observation)
        row["observation_date"] = observation.observation_date.isoformat()
        writer.writerow({key: row.get(key, "") if row.get(key) is not None else "" for key in headers})
    return send_file(io.BytesIO(output.getvalue().encode("utf-8-sig")), mimetype="text/csv", as_attachment=True, download_name=f"crop_cycle_{cycle.id}_observations.csv")
